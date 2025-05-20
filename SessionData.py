import main
import os
import datetime as dt
from openpyxl import load_workbook

class SessionData: # Sesijas dati
    def __init__(self):
        self.entries = {} # izveidojam tukšu vārdnīcu
        self.total_duration = 0.0
    def load_from_excel(self, filename): # ielādējam datus no Excel faila, ja tāds eksistē	
        if not os.path.exists(filename):
            return
        wb = load_workbook(filename)
        ws = wb.active

        for row in ws.iter_rows(min_row=2):
            try:
                start_time = dt.datetime.strptime(row[3].value, "%Y-%m-%d %H:%M:%S")
                end_time = dt.datetime.strptime(row[4].value, "%Y-%m-%d %H:%M:%S")
                entry = {
                    "title": row[0].value,
                    "process": row[1].value,
                    "duration": self._hhmm_to_seconds(row[2].value),
                    "start": start_time,
                    "end": dt.datetime.now(),
                }
                key = f"{entry['process']} - {entry['title']}"
                self.entries[key] = entry
            except Exception as e:
                print(f"Error loading entry from Excel: {e}")

    def _hhmm_to_seconds(self, hhmm): # konvertējam HH:MM formātu uz sekundēm
        hours, minutes = map(int, hhmm.split(':'))
        return hours * 3600 + minutes * 60
    
    def add_entry(self, key, process, title, start_time): # pievienojam jaunu ierakstu
        self.entries[key] = {
            "process": process,
            "title": title,
            "start": start_time,
            "end": start_time,
            "duration": 0.0
        }
    def update_entry(self, key, end_time): # atjaunojam esošo ierakstu
        entry = self.entries[key]
        duration = (end_time - entry["end"]).total_seconds()
        self.total_duration += entry["duration"]
        entry["duration"] += duration
        entry["end"] = end_time
        self.total_duration += duration
    def filter_entries(self, min_duration=1): # filtrējam ierakstus
        return {k: v for k, v in self.entries.items() if v["duration"] >= min_duration and v["process"] not in main.SYSTEM_PROCESSES}
    def get_sorted_entries(self, sort_key="process"): # iegūstam sakārtotus ierakstus pēc atslēgas
        filtered = self.filter_entries()
        return sorted(filtered.items(), key=lambda x: x[1][sort_key])