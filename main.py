import win32gui as win
import win32process as proc
import time
import datetime as dt
import psutil 
from SessionData import SessionData
from openpyxl import Workbook
from openpyxl.styles import PatternFill
import matplotlib.pyplot as plt


SYSTEM_PROCESSES = {'Explorer', 'Svchost', 'Taskmgr'}
CATEGORIES = {
    "Browser": {
        "processes": ["Chrome", "Firefox", "msedge", "opera"],
        "color": "FFB8F0B4"  
    },
    "Work": {
        "processes": ["Code", "Pycharm", "Devenv", "Excel", "Winword"],
        "color": "FFF0E68C"
    },
    "Games": {
        "processes": ["Steam", "Dota2", "Cs2", "Game"],
        "color": "FFFFC0CB"
    },
    "Messenger": {
        "processes": ["Telegram", "Discord", "Whatsapp"],
        "color": "FFADD8E6"
    },
}
DEFAULT_CATEGORY = {
    "name": "Other",
    "color": "FFFFFFFF"
}

def get_active_window(): # iegūstam aktīvo logu (nosaukumu un procesu)
    try:
        window = win.GetForegroundWindow()
        pid = proc.GetWindowThreadProcessId(window)[1]
        exetubale = psutil.Process(pid).name().capitalize().removesuffix(".exe")
        title = win.GetWindowText(window)
        return {"process": exetubale, "title": title.strip()}
    except Exception as e:
        print(f"Error getting active window: {e}")
        return None

def track_activity():
    session = SessionData()
    current_window = None
    print("Sessija tiek ierakstīta")
    try:
        while True:
            new_window = get_active_window()
            if not new_window:
                continue

            new_key = f"{new_window['process']} - {new_window['title']}"
            if new_key != current_window:
                if current_window:
                    session.update_entry(current_window, dt.datetime.now())

                if not any(e.process == new_window['process'] and e.title == new_window['title']
                           for e in session._entries):
                    session.add_entry(key=new_key, process=new_window['process'],
                                      title=new_window['title'], start_time=dt.datetime.now())
                current_window = new_key

            time.sleep(1)

    except KeyboardInterrupt:
        print("\nSessijas beigas...")
        if current_window:
            session.update_entry(current_window, dt.datetime.now())

        # atkļūdošanai
        print("Session entries:")
        for entry in session._entries:
            print(f"{entry.process} - {entry.title}: {entry.duration}s")


        _save_to_excel(session)
        analyze_session(session)


def _save_to_excel(data: SessionData): # ieraksta saglabāšana Excel failā
    today = dt.datetime.now().strftime("%Y-%m-%d")
    filename = f"session_data_{today}.xlsx"

    try:
        wb = Workbook()
        ws = wb.active
        ws.title = "Session Report"
        headers = ["Window Title", "Process", "Duration (HH:MM)", "Start Time", "End Time", "Activity type"]
        ws.append(headers)

        styles = {cat: PatternFill(start_color=data["color"], fill_type="solid")
                  for cat, data in CATEGORIES.items()}
        default_style = PatternFill(start_color=DEFAULT_CATEGORY["color"], fill_type="solid")

        for key, entry in data.get_sorted_entries(sort_key="process"):
            category = get_category(entry["process"])
            duration = seconds_to_hhmm(entry["duration"])

            ws.append([
                entry["title"],
                entry["process"],
                duration,
                entry["start"].strftime("%Y-%m-%d %H:%M:%S"),
                entry["end"].strftime("%Y-%m-%d %H:%M:%S"),
                category["name"]
            ])

            ws.cell(row=ws.max_row, column=6).fill = styles.get(category["name"], default_style)

        wb.save(filename)
        print(f"Dati saglabāti: {filename}")
        return True

    except Exception as e:
        print(f"Kļūda saglābājot Excel: {e}")
        return False

def seconds_to_hhmm(seconds):
    hours = seconds // 3600 # 3600 sekundes = 1 stunda (saglabam tikati veselu skaitli)
    minutes = (seconds % 3600) // 60 # 60 sekundes = 1 minūte
    return f"{int(hours):02}:{int(minutes):02}"

def get_category(process_name): # piešķiram kategoriju katram procesam
    for category, data in CATEGORIES.items():
        if any(process_name in p for p in data["processes"]):
            return {
                "name": category,
                "color": data["color"]
            }
    return DEFAULT_CATEGORY


def analyze_session(session): # datu analīze un diagrammas konstruēšana
    if not session._entries:
        print("Nav datu analīzei")
        return

    longest_entry = max(session._entries, key=lambda x: x.duration) # ilgāka loga atrašana
    print(f"\nСАМОЕ ДОЛГОЕ ОКНО:")
    print(f"Процесс: {longest_entry.process}")
    print(f"Заголовок: {longest_entry.title}")
    print(f"Время: {seconds_to_hhmm(longest_entry.duration)}")

    categories = {} #datu savākšana
    for entry in session._entries:
        category = get_category(entry.process)['name']
        categories[category] = categories.get(category, 0) + entry.duration


    valid_categories = {} # nulles vērtību filtrēšana un krāsu pārveidošana
    color_map = {}

    for cat, duration in categories.items():
        if duration > 0:
            valid_categories[cat] = duration
            hex_color = CATEGORIES.get(cat, DEFAULT_CATEGORY)['color'] # pārveidojam HEX krāsu formātu matplotlib saprotamajā formātā
            color_map[cat] = f'#{hex_color[2:]}'  # Pārveidojam "FFB8F0B4" -> "#B8F0B4"

    if not valid_categories:
        print("\nNav datu diagrammai")
        return

    plt.figure(figsize=(10, 6)) # diagrammas sastādīšana
    labels = list(valid_categories.keys())
    sizes = list(valid_categories.values())
    colors = [color_map[label] for label in labels]

    wedgeprops = {
        'linewidth': 1,  # Kontūra treknums
        'edgecolor': 'black'  # Kontūra krāsa (var mainīt uz 'white' vai 'gray')
    }

    plt.pie(sizes, labels=labels, colors=colors, autopct='%1.1f%%', startangle=140,wedgeprops=wedgeprops)
    plt.axis('equal')
    plt.title('Atskaite')

    plt.legend(labels, title="Kategorijas", loc="best") # leģenda

    plt.show()

if __name__ == "__main__":
    track_activity()

